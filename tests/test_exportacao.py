import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from pydosap.services.exportacao import exportar_ctes_excel


class ExportarCtesExcelTests(unittest.TestCase):
    def test_gera_planilha_com_cabecalho_e_formatos(self):
        colunas = ["lancado", "regiao", "dataEmissao", "valor", "resposta"]
        linhas = [
            {"lancado": True, "regiao": "NORDESTE", "dataEmissao": "2026-09-18T00:00:00",
             "valor": 1234.5, "resposta": "OK"},
            {"lancado": False, "regiao": "SUDESTE", "dataEmissao": "2026-09-01",
             "valor": 99.9, "resposta": ""},
        ]

        with tempfile.TemporaryDirectory() as tmp:
            caminho = Path(tmp) / "relatorio.xlsx"
            exportar_ctes_excel(str(caminho), colunas, linhas)

            wb = load_workbook(caminho)
            ws = wb.active

            self.assertEqual(
                [c.value for c in ws[1]], ["Lançado", "Região", "Emissão", "Valor", "Resposta"],
            )
            self.assertEqual(ws.cell(row=2, column=1).value, "Sim")
            self.assertEqual(ws.cell(row=3, column=1).value, "Não")
            self.assertEqual(ws.cell(row=2, column=3).value, "18/09/2026")
            self.assertEqual(ws.cell(row=2, column=4).value, 1234.5)
            self.assertEqual(ws.cell(row=2, column=4).number_format, '"R$" #,##0.00')
            self.assertEqual(ws.freeze_panes, "A2")

    def test_linhas_vazias_nao_quebra(self):
        with tempfile.TemporaryDirectory() as tmp:
            caminho = Path(tmp) / "vazio.xlsx"
            exportar_ctes_excel(str(caminho), ["chave"], [])
            wb = load_workbook(caminho)
            self.assertEqual(wb.active.cell(row=1, column=1).value, "Chave")


if __name__ == "__main__":
    unittest.main()
