"""Exportação da tabela de CT-es de despesa para uma planilha do Excel."""
from __future__ import annotations

from ..cte_columns import COLUNAS_DATA, COLUNAS_MOEDA, rotulo
from ..formatting import data_br

FONTE = "Arial"
COR_CABECALHO_FUNDO = "212327"   # BG_2 do tema escuro (ui/styles.py)
COR_CABECALHO_TEXTO = "E8E9EB"   # TEXT do tema escuro
LARGURA_MINIMA, LARGURA_MAXIMA = 8, 60


def exportar_ctes_excel(caminho: str, colunas: list[str], linhas: list[dict]) -> None:
    """Grava `linhas` (na ordem de `colunas`) num .xlsx, com cabeçalho destacado,
    coluna de data em dd/mm/aaaa e coluna de valor em formato contábil BRL.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "CT-es"

    fonte_cabecalho = Font(name=FONTE, bold=True, color=COR_CABECALHO_TEXTO)
    fundo_cabecalho = PatternFill("solid", fgColor=COR_CABECALHO_FUNDO)
    fonte_corpo = Font(name=FONTE)

    for j, coluna in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=j, value=rotulo(coluna))
        cel.font = fonte_cabecalho
        cel.fill = fundo_cabecalho
        cel.alignment = Alignment(horizontal="center", vertical="center")

    for i, linha in enumerate(linhas, start=2):
        for j, coluna in enumerate(colunas, start=1):
            valor = linha.get(coluna)
            cel = ws.cell(row=i, column=j)
            cel.font = fonte_corpo
            if coluna in COLUNAS_DATA:
                cel.value = data_br(valor)
            elif coluna in COLUNAS_MOEDA:
                try:
                    cel.value = float(valor) if valor not in (None, "") else None
                except (TypeError, ValueError):
                    cel.value = valor
                cel.number_format = '"R$" #,##0.00'
            elif isinstance(valor, bool):
                cel.value = "Sim" if valor else "Não"
            else:
                cel.value = valor

    for j, coluna in enumerate(colunas, start=1):
        maior = len(rotulo(coluna))
        for linha in linhas:
            maior = max(maior, len(str(linha.get(coluna) or "")))
        ws.column_dimensions[get_column_letter(j)].width = min(
            max(maior + 2, LARGURA_MINIMA), LARGURA_MAXIMA
        )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(caminho)
